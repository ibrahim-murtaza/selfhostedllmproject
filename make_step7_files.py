"""
Makes three fictional, non-confidential test files for Phase C step 7:

  scanned.pdf  2 pages, image only (NO text layer) -> forces OCR
  test.docx    Word file with headings, bullets and a table
  tables.pdf   40-row table -> exercises Docling's table extraction

Run:
    pip install python-docx
    python make_step7_files.py

The expected answers are printed at the end.
"""

import os
import random

from docx import Document
from PIL import Image, ImageDraw, ImageFilter, ImageFont
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.pdfgen import canvas
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

random.seed(7)


# ---------------------------------------------------------------- scanned.pdf
def load_font(size: int):
    for name in ("arial.ttf", "Arial.ttf", "DejaVuSans.ttf", "LiberationSans-Regular.ttf"):
        try:
            return ImageFont.truetype(name, size)
        except OSError:
            continue
    return ImageFont.load_default(size=size)  # Pillow >= 10.1


SCAN_PAGES = [
    [
        "Scanned Contract Summary (fictional)",
        "",
        "Client: Contoso Test Ltd",
        "Effective date: 14 March 2026",
        "Payment terms: Net 45 days",
        "Support tier: Standard",
        "",
        "Scan code: SCAN-1-CODE-4821",
    ],
    [
        "Schedule B - Termination and Law",
        "",
        "Termination notice: 60 days",
        "Governing law: Ontario",
        "Renewal: automatic, 12 months",
        "",
        "Scan code: SCAN-2-CODE-7305",
    ],
]


def render_scan_page(lines):
    w, h = 1240, 1754  # A4 at 150 dpi
    img = Image.new("L", (w, h), 250)
    d = ImageDraw.Draw(img)
    title_font, body_font = load_font(46), load_font(34)
    y = 160
    for i, line in enumerate(lines):
        d.text((110, y), line, font=title_font if i == 0 else body_font, fill=25)
        y += 90 if i == 0 else 62
    # mild scan-like imperfections
    px = img.load()
    for _ in range(9000):
        x, yy = random.randrange(w), random.randrange(h)
        px[x, yy] = max(0, px[x, yy] - random.randint(20, 70))
    img = img.rotate(0.4, resample=Image.BICUBIC, fillcolor=250)
    return img.filter(ImageFilter.GaussianBlur(0.5))


def make_scanned(path="scanned.pdf"):
    c = canvas.Canvas(path, pagesize=A4)
    for i, lines in enumerate(SCAN_PAGES):
        tmp = f"_scan_page_{i}.png"
        render_scan_page(lines).save(tmp)
        c.drawImage(tmp, 0, 0, width=A4[0], height=A4[1])
        c.showPage()
        os.remove(tmp)
    c.save()


# ------------------------------------------------------------------ test.docx
def make_docx(path="test.docx"):
    doc = Document()
    doc.add_heading("Northwind Test Corp - Onboarding Checklist (Sample)", level=0)
    doc.add_heading("1. Overview", level=1)
    doc.add_paragraph(
        "This fictional checklist covers the first week for a new hire. "
        "Reference code: DOCX-ALPHA-5512."
    )
    doc.add_heading("2. First-week tasks", level=1)
    for item in (
        "Collect laptop and badge from the IT desk",
        "Complete the security awareness module",
        "Meet your manager for a 30-minute intro",
    ):
        doc.add_paragraph(item, style="List Bullet")
    doc.add_heading("3. Owners and deadlines", level=1)
    table = doc.add_table(rows=1, cols=3)
    table.style = "Light Grid Accent 1"
    for cell, text in zip(table.rows[0].cells, ("Task", "Owner", "Due")):
        cell.text = text
    for row in (
        ("Laptop setup", "IT desk", "Day 1"),
        ("Security module", "New hire", "Day 3"),
        ("Manager intro", "Line manager", "Day 2"),
    ):
        cells = table.add_row().cells
        for cell, text in zip(cells, row):
            cell.text = text
    doc.add_heading("4. Contact", level=1)
    doc.add_paragraph("Questions go to People Operations. Closing code: DOCX-OMEGA-8830.")
    doc.save(path)


# ------------------------------------------------------------------ tables.pdf
def q_units(n: int, q: int) -> int:
    return (n * (30 + q * 7)) % 900 + 100  # deterministic, looks arbitrary


def make_tables(path="tables.pdf"):
    styles = getSampleStyleSheet()
    data = [["SKU", "Q1 units", "Q2 units", "Q3 units", "Q4 units", "Unit price (USD)"]]
    for n in range(1, 41):
        data.append(
            [f"SKU-{n:03d}"]
            + [str(q_units(n, q)) for q in (1, 2, 3, 4)]
            + [f"{5 + (n * 13) % 95}.{(n * 7) % 100:02d}"]
        )
    t = Table(data, colWidths=[3 * cm, 2.8 * cm, 2.8 * cm, 2.8 * cm, 2.8 * cm, 3.6 * cm], repeatRows=1)
    t.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0F6E6E")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F1F5F5")]),
                ("ALIGN", (1, 1), (-1, -1), "RIGHT"),
                ("PADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    SimpleDocTemplate(path, pagesize=A4, topMargin=2 * cm, bottomMargin=2 * cm).build(
        [
            Paragraph("Fictional Unit Sales by SKU", styles["Title"]),
            Paragraph("All figures are made up for testing.", styles["Normal"]),
            Spacer(1, 10),
            t,
        ]
    )


make_scanned()
make_docx()
make_tables()

print("Wrote scanned.pdf, test.docx, tables.pdf\n")
print("Expected answers:")
print("  scanned.pdf : 'what are the payment terms?'          -> Net 45 days")
print("                'what is the scan code on page 2?'     -> SCAN-2-CODE-7305")
print("  test.docx   : 'what is the closing code?'            -> DOCX-OMEGA-8830")
print("                'who owns the security module task?'   -> New hire (due Day 3)")
print(f"  tables.pdf  : 'what are the Q3 units for SKU-027?'   -> {q_units(27, 3)}")
print(f"                'what are the Q1 units for SKU-005?'   -> {q_units(5, 1)}")