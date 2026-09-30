import io
import sys
import time

import numpy as np
from PIL import Image
from pptx import Presentation
from rapidocr import RapidOCR

prs = Presentation(sys.argv[1] if len(sys.argv) > 1 else "test.pptx")
engine = RapidOCR()

for n, slide in enumerate(prs.slides, 1):
    for sh in slide.shapes:
        if not hasattr(sh, "image"):
            continue
        img = np.array(Image.open(io.BytesIO(sh.image.blob)).convert("RGB"))
        for attempt in (1, 2):
            t0 = time.monotonic()
            out = engine(img)
            dt = time.monotonic() - t0
            txts = getattr(out, "txts", None)
            print(f"[slide {n}] run {attempt}: {dt:.2f}s -> {list(txts) if txts else txts}")