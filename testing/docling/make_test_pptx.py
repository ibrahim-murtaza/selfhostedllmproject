"""
Makes a fictional, non-confidential 6-slide test deck and a facts file:

  test.pptx              the deck
  test_pptx_facts.json   what a good parser should find (used by compare_pptx.py)

Each slide tests one thing parsers often get wrong:
  1 title + subtitle            2 nested bullets          3 a table
  4 a native chart (values live in embedded data, not on the slide)
  5 a picture that contains text (needs OCR)              6 speaker notes

Run:
    pip install python-pptx
    python make_test_pptx.py
"""

import json

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches, Pt

prs = Presentation()  # default 4:3 template

# ---- 1. title slide
s = prs.slides.add_slide(prs.slide_layouts[0])
s.shapes.title.text = "Northwind Roadmap Review"
s.placeholders[1].text = "Prepared by the Test Team - code PPTX-ALPHA-3141"

# ---- 2. nested bullets
s = prs.slides.add_slide(prs.slide_layouts[1])
s.shapes.title.text = "Priorities"
tf = s.placeholders[1].text_frame
tf.text = "Ship the mobile app"
for text, level in (
    ("Pilot with two departments", 1),
    ("Cut onboarding time to 5 days", 0),
    ("Consolidate vendor contracts", 0),
    ("Confirm the shortlist by June", 1),
):
    p = tf.add_paragraph()
    p.text = text
    p.level = level

# ---- 3. table
s = prs.slides.add_slide(prs.slide_layouts[5])
s.shapes.title.text = "Regional Sales"
rows = (
    ("Region", "Q1", "Q2", "Q3", "Q4"),
    ("North", 1200, 1350, 1100, 1500),
    ("South", 900, 980, 1020, 1100),
    ("East", 1450, 1600, 1580, 1720),
    ("West", 800, 875, 910, 990),
    ("Central", 1010, 1090, 1150, 1230),
)
table = s.shapes.add_table(len(rows), 5, Inches(0.6), Inches(1.8), Inches(8.8), Inches(3.6)).table
for r, row in enumerate(rows):
    for c, value in enumerate(row):
        table.cell(r, c).text = str(value)

# ---- 4. native chart
s = prs.slides.add_slide(prs.slide_layouts[5])
s.shapes.title.text = "Team Size"
chart_data = CategoryChartData()
chart_data.categories = ["Ops", "Finance", "Support", "Engineering"]
chart_data.add_series("Headcount by Team", (26, 19, 33, 47))
gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.8), Inches(1.8), Inches(8.4), Inches(4.8), chart_data)
gf.chart.has_title = True
gf.chart.chart_title.text_frame.text = "Headcount by Team"

# ---- 5. picture containing text
try:
    font = ImageFont.truetype("arial.ttf", 54)
except OSError:
    try:
        font = ImageFont.truetype("DejaVuSans.ttf", 54)
    except OSError:
        font = ImageFont.load_default(size=54)
img = Image.new("RGB", (1200, 420), (238, 244, 244))
d = ImageDraw.Draw(img)
d.rectangle([12, 12, 1188, 408], outline=(15, 110, 110), width=6)
d.text((70, 90), "Site inspection photo", font=font, fill=(20, 20, 20))
d.text((70, 230), "Scan code: IMG-7702", font=font, fill=(20, 20, 20))
img.save("_slide5.png")
s = prs.slides.add_slide(prs.slide_layouts[5])
s.shapes.title.text = "Site Photo"
s.shapes.add_picture("_slide5.png", Inches(0.8), Inches(2.0), width=Inches(8.4))

# ---- 6. bullets + speaker notes
s = prs.slides.add_slide(prs.slide_layouts[1])
s.shapes.title.text = "Next Steps"
tf = s.placeholders[1].text_frame
tf.text = "Confirm budget with Finance"
p = tf.add_paragraph()
p.text = "Book the launch review"
s.notes_slide.notes_text_frame.text = (
    "Speaker notes: the launch date is 14 March 2027. Internal code PPTX-OMEGA-9926."
)

prs.save("test.pptx")

import os

os.remove("_slide5.png")

# ---------------------------------------------------------------- facts file
# "re" facts are regular expressions (case-insensitive, . matches newlines).
facts = [
    {"id": "title", "slide": 1, "kind": "title", "re": r"Northwind Roadmap Review"},
    {"id": "subtitle_code", "slide": 1, "kind": "subtitle", "re": r"PPTX-ALPHA-3141"},
    {"id": "bullet_top", "slide": 2, "kind": "bullet", "re": r"Ship the mobile app"},
    {"id": "bullet_nested", "slide": 2, "kind": "nested bullet", "re": r"Pilot with two departments"},
    {"id": "bullet_priority2", "slide": 2, "kind": "bullet", "re": r"Cut onboarding time to 5 days"},
    {"id": "table_header", "slide": 3, "kind": "table header", "re": r"Region.{0,20}Q1.{0,20}Q2"},
    {"id": "table_west_q2", "slide": 3, "kind": "table cell in the right row", "re": r"West[^A-Za-z]{0,60}875"},
    {"id": "table_east_q4", "slide": 3, "kind": "table cell in the right row", "re": r"East[^A-Za-z]{0,60}1720"},
    {"id": "chart_title", "slide": 4, "kind": "chart title", "re": r"Headcount by Team"},
    {"id": "chart_engineering", "slide": 4, "kind": "chart data value", "re": r"Engineering[^0-9]{0,30}47"},
    {"id": "image_text", "slide": 5, "kind": "text inside a picture", "re": r"IMG-7702"},
    {"id": "notes_date", "slide": 6, "kind": "speaker notes", "re": r"14 March 2027"},
    {"id": "notes_code", "slide": 6, "kind": "speaker notes", "re": r"PPTX-OMEGA-9926"},
    {"id": "bullet_last", "slide": 6, "kind": "bullet", "re": r"Confirm budget with Finance"},
]
qa = [
    {"q": "What is the code on the title slide?", "any": ["PPTX-ALPHA-3141"]},
    {"q": "What is the second priority on the Priorities slide?", "any": ["onboarding"]},
    {"q": "What were the West region's Q2 sales?", "any": ["875"]},
    {"q": "How many people are on the Engineering team according to the chart?", "any": ["47"]},
    {"q": "What code is written in the picture on the Site Photo slide?", "any": ["IMG-7702"]},
    {"q": "What launch date appears in the speaker notes?", "any": ["14 March 2027"]},
    {"q": "How many slides does the presentation have?", "any": ["6", "six"]},
]
with open("test_pptx_facts.json", "w", encoding="utf-8") as f:
    json.dump({"file": "test.pptx", "slides": 6, "facts": facts, "qa": qa}, f, indent=2)

print("Wrote test.pptx (6 slides) and test_pptx_facts.json")
print("Slides expected to be hard: 4 (chart values), 5 (text inside a picture), 6 (speaker notes).")
