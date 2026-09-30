import io
import sys
import time

from docling.datamodel.base_models import DocumentStream
from docling.document_converter import DocumentConverter
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

prs = Presentation(sys.argv[1] if len(sys.argv) > 1 else "test.pptx")
conv = DocumentConverter()


def pictures(shapes):
    for sh in shapes:
        try:
            kind = sh.shape_type
        except Exception:
            continue
        if kind == MSO_SHAPE_TYPE.GROUP:
            yield from pictures(sh.shapes)
        elif kind == MSO_SHAPE_TYPE.PICTURE:
            yield sh


for n, slide in enumerate(prs.slides, 1):
    if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
        text = slide.notes_slide.notes_text_frame.text.strip()
        if text:
            print(f"[slide {n}] NOTES: {text}")
    for sh in pictures(slide.shapes):
        img = sh.image
        t0 = time.monotonic()
        res = conv.convert(
            DocumentStream(name=f"slide{n}.{img.ext}", stream=io.BytesIO(img.blob))
        )
        md = res.document.export_to_markdown().strip()
        print(
            f"[slide {n}] PICTURE {img.ext} {len(img.blob)} bytes, "
            f"{time.monotonic() - t0:.1f}s: {md!r}"
        )