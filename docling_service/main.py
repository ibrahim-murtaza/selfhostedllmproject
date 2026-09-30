"""
Docling conversion service -- CPU-only, persistent, localhost-only.

POST /convert   raw file bytes in the body, filename in the X-Filename header
                -> {"markdown", "chars", "cached", "convert_ms"}
GET  /health    -> {"status": "ok", "cache_entries": n}

Design notes:
- CPU-only: GPU is hidden from torch, and the PDF pipeline is pinned to CPU.
  (Docling's GPU backend leaks VRAM in a long-running process.)
- Models are loaded at startup, so the first user request is not slow.
- Page limit (MAX_PAGES): Docling checks the page count when it opens the
  file, before any model runs, so an over-long PDF is refused in milliseconds
  instead of after minutes of CPU. Conversion costs roughly 3-4 s per page.
- One conversion at a time (lock). Pilot scale; also makes the cache safe:
  a second identical request waits, then gets a cache hit.
- Results are cached by sha256 of the file bytes (bounded, oldest evicted),
  because the gateway sees the same file again on every chat turn.
- No auth: bind to 127.0.0.1 only. Do not expose this port.

Run:
    uvicorn main:app --host 127.0.0.1 --port 8001
"""

import hashlib
import os
import re
import threading
import time
from collections import OrderedDict
from contextlib import asynccontextmanager
from io import BytesIO

# Hide GPUs before torch is imported anywhere.
os.environ["CUDA_VISIBLE_DEVICES"] = ""

from docling.datamodel.base_models import DocumentStream, InputFormat
from docling.datamodel.pipeline_options import (
    AcceleratorDevice,
    AcceleratorOptions,
    PdfPipelineOptions,
)
from docling.document_converter import DocumentConverter, PdfFormatOption
from docling.exceptions import ConversionError
from fastapi import FastAPI, HTTPException, Request
from fastapi.responses import JSONResponse
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from starlette.concurrency import run_in_threadpool

MAX_BYTES = 25 * 1024 * 1024  # reject anything larger
# ~28 pages of dense text is about all the Quick model's window can hold
# (see the gateway's size guard); 30 keeps a little headroom and caps the
# worst-case conversion time at roughly two minutes.
MAX_PAGES = 30
CACHE_MAX = 32  # documents kept in memory
MAX_PICTURES = 10  # pictures read (OCR) per PowerPoint file
MIN_PICTURE_PX = 120  # skip icons and logos smaller than this (width or height)

_pipeline_options = PdfPipelineOptions()
_pipeline_options.accelerator_options = AcceleratorOptions(device=AcceleratorDevice.CPU)

_converter = DocumentConverter(
    format_options={InputFormat.PDF: PdfFormatOption(pipeline_options=_pipeline_options)}
)

_lock = threading.Lock()
_cache: "OrderedDict[str, str]" = OrderedDict()


@asynccontextmanager
async def lifespan(app: FastAPI):
    # Load the PDF models now instead of on the first user request.
    t0 = time.monotonic()
    try:
        _converter.initialize_pipeline(InputFormat.PDF)
        print(f"[docling] models warmed in {time.monotonic() - t0:.1f}s")
    except Exception as e:  # warm-up is an optimisation, never a blocker
        print(f"[docling] warm-up skipped ({type(e).__name__}: {e})")
    yield


app = FastAPI(lifespan=lifespan)


@app.exception_handler(HTTPException)
async def openai_style_errors(request: Request, exc: HTTPException):
    return JSONResponse(
        status_code=exc.status_code,
        content={"error": {"message": exc.detail, "code": exc.status_code}},
    )


def _clean(markdown: str) -> str:
    # Collapse runs of spaces inside a line (justified-text artefact).
    # Leading indentation is left alone.
    return re.sub(r"(?<=\S)[ \t]{2,}(?=\S)", " ", markdown)


_PICTURE_EXTS = {"png", "jpg", "jpeg", "bmp", "tif", "tiff", "webp"}


def _pictures(shapes):
    """Every picture on a slide, including ones inside groups."""
    for sh in shapes:
        try:
            if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
                yield from _pictures(sh.shapes)
            elif hasattr(sh, "image"):
                yield sh
        except Exception:
            continue


_ocr_engine = None


def _ocr_picture(blob: bytes) -> str:
    """Text found in one picture. Calls RapidOCR (already installed for
    Docling) directly: far faster than running a whole Docling conversion on
    the image. The engine loads on first use; the caller holds _lock."""
    global _ocr_engine
    import numpy as np
    from PIL import Image
    from rapidocr import RapidOCR

    if _ocr_engine is None:
        _ocr_engine = RapidOCR()
    out = _ocr_engine(np.array(Image.open(BytesIO(blob)).convert("RGB")))
    return " ".join(out.txts) if getattr(out, "txts", None) else ""


def _pptx_extras(data: bytes) -> str:
    """Speaker notes and text inside pictures: the two things Docling's
    PowerPoint reader leaves out. Returns Markdown to append to Docling's
    output ("" if there is nothing). Never raises: a failure here must not
    break the conversion. The caller holds _lock (this uses the shared OCR engine)."""
    try:
        prs = Presentation(BytesIO(data))
    except Exception:
        return ""
    notes: list[str] = []
    pictures: list[str] = []
    seen: set[str] = set()
    for n, slide in enumerate(prs.slides, 1):
        try:
            frame = slide.notes_slide.notes_text_frame if slide.has_notes_slide else None
            if frame is not None and frame.text.strip():
                notes.append(f"Slide {n}: {frame.text.strip()}")
            for sh in _pictures(slide.shapes):
                if len(seen) >= MAX_PICTURES:
                    break
                try:
                    img = sh.image
                    width, height = img.size
                    if img.ext not in _PICTURE_EXTS or min(width, height) < MIN_PICTURE_PX:
                        continue
                    digest = hashlib.sha1(img.blob).hexdigest()
                    if digest in seen:  # same logo on every slide: read it once
                        continue
                    seen.add(digest)
                    text = " ".join(_ocr_picture(img.blob).split())
                    if text:
                        pictures.append(f"Slide {n}: {text}")
                except Exception:
                    continue
        except Exception:
            continue
    out = ""
    if notes:
        out += "\n\n## Speaker notes\n\n" + "\n\n".join(notes)
    if pictures:
        out += "\n\n## Text found in pictures\n\n" + "\n\n".join(pictures)
    return out


def _convert(name: str, data: bytes, key: str) -> tuple[str, bool, float]:
    with _lock:
        if key in _cache:
            _cache.move_to_end(key)
            return _cache[key], True, 0.0
        t0 = time.monotonic()
        result = _converter.convert(
            DocumentStream(name=name, stream=BytesIO(data)),
            max_num_pages=MAX_PAGES,
        )
        markdown = _clean(result.document.export_to_markdown())
        if name.lower().endswith(".pptx"):
            markdown += _pptx_extras(data)
        _cache[key] = markdown
        while len(_cache) > CACHE_MAX:
            _cache.popitem(last=False)
        return markdown, False, (time.monotonic() - t0) * 1000


@app.post("/convert")
async def convert(request: Request):
    data = await request.body()
    if not data:
        raise HTTPException(status_code=400, detail="Empty request body.")
    if len(data) > MAX_BYTES:
        raise HTTPException(
            status_code=413,
            detail=f"File too large ({len(data) // (1024 * 1024)} MB, limit {MAX_BYTES // (1024 * 1024)} MB).",
        )
    name = request.headers.get("x-filename", "document.pdf")
    key = hashlib.sha256(data).hexdigest()
    try:
        markdown, cached, ms = await run_in_threadpool(_convert, name, data, key)
    except ConversionError as e:
        # Docling's message is long and includes internal paths; keep only the
        # reason, and say the page limit in plain words.
        reason = str(e).split("Errors: ", 1)[-1]
        m = re.search(r"Document has (\d+) pages", reason)
        if m:
            reason = f"Document has {m.group(1)} pages; the limit is {MAX_PAGES}."
        # No file name in the message: the gateway adds it ("Could not read 'x': ...").
        raise HTTPException(status_code=422, detail=reason)
    except Exception as e:
        raise HTTPException(status_code=422, detail=f"{type(e).__name__}: {e}")
    return {
        "markdown": markdown,
        "chars": len(markdown),
        "cached": cached,
        "convert_ms": round(ms, 1),
    }


@app.get("/health")
def health():
    return {"status": "ok", "cache_entries": len(_cache), "max_pages": MAX_PAGES}